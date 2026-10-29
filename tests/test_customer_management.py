from datetime import date
from io import BytesIO
import re
from unittest.mock import MagicMock
import unittest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from customer_data import SCHEMAS, CustomerStore, normalize_record, upcoming_tanks
from customer_excel import workbook_bytes, read_workbook


def customer(code='C001',name='테스트 거래처'):
 return normalize_record('거래처DB',{'거래처코드':code,'거래처명':name,'거래구분':'직거래처'},set())


def fake_service(tables):
 service=MagicMock()
 def metadata(**kwargs):
  req=MagicMock();req.execute.return_value={'sheets':[{'properties':{'title':s}} for s in tables]};return req
 service.spreadsheets.return_value.get.side_effect=metadata
 def read(**kwargs):
  sheet=kwargs['range'].split("'")[1]
  req=MagicMock();req.execute.return_value={'values':tables[sheet]};return req
 service.spreadsheets.return_value.values.return_value.get.side_effect=read
 def creates(**kwargs):
  for r in kwargs['body']['requests']:tables[r['addSheet']['properties']['title']]=[]
  req=MagicMock();req.execute.return_value={};return req
 service.spreadsheets.return_value.batchUpdate.side_effect=creates
 def write(**kwargs):
  for block in kwargs['body']['data']:
   match=re.match(r"'(.+)'!A(\d+)",block['range']);sheet,start=match[1],int(match[2])-1
   rows=tables[sheet]
   while len(rows)<start+len(block['values']):rows.append([])
   rows[start:start+len(block['values'])]=block['values']
  req=MagicMock();req.execute.return_value={};return req
 service.spreadsheets.return_value.values.return_value.batchUpdate.side_effect=write
 return service


class CustomerManagementTests(unittest.TestCase):
 def test_excel_template_is_styled_blank_and_roundtrips(self):
  data=workbook_bytes();parsed=read_workbook(data)
  self.assertEqual(set(parsed),set(SCHEMAS))
  self.assertTrue(all(not rows for rows in parsed.values()))
  wb=load_workbook(BytesIO(data))
  self.assertEqual(wb['거래처DB'].freeze_panes,'C2')
  self.assertTrue(wb['거래처DB'].data_validations.dataValidation)
  self.assertEqual(wb['거래처DB']['A2'].number_format,'@')
  self.assertIn('항목설명',wb.sheetnames)
  wb.close()

 def test_excel_rejects_formulas_and_header_changes(self):
  for mode in ('formula','header'):
   wb=load_workbook(BytesIO(workbook_bytes()))
   wb['거래처DB']['A2' if mode=='formula' else 'A1']='=1+1' if mode=='formula' else 'changed'
   buffer=BytesIO();wb.save(buffer);wb.close()
   with self.assertRaises(ValueError):read_workbook(buffer.getvalue())

 def test_filled_workbook_import_preserves_phone_and_customer_links(self):
  wb=load_workbook(BytesIO(workbook_bytes()))
  client=customer()
  for col,field in enumerate(SCHEMAS['거래처DB'],1):wb['거래처DB'].cell(2,col,client[field])
  row={'거래처코드':'C001','성명':'구매 담당자','연락처':'010-0000-0000'}
  for col,field in enumerate(SCHEMAS['거래처연락처DB'],1):wb['거래처연락처DB'].cell(2,col,row.get(field,''))
  buffer=BytesIO();wb.save(buffer);wb.close()
  incoming=read_workbook(buffer.getvalue())
  tables={};store=CustomerStore(fake_service(tables),'sid')
  self.assertEqual(store.import_new(incoming),2)
  self.assertEqual(store.read_all()['거래처연락처DB'][0]['연락처'],'010-0000-0000')

 def test_validation_and_links(self):
  contact=normalize_record('거래처연락처DB',{'거래처코드':'C001','성명':'담당자','연락처':'010-1234-5678'},{'C001'})
  self.assertTrue(contact['연락처ID'])
  self.assertEqual(contact['연락처'],'010-1234-5678')
  with self.assertRaises(ValueError):normalize_record('거래처연락처DB',{'거래처코드':'missing','성명':'담당자'},set())
  for bad in ({'용량':'-1'},{'월임대료':'NaN'},{'계약시작일':'2026-02-30'},{'계약시작일':'2026-11-01','계약종료일':'2026-10-01'}):
   with self.subTest(bad=bad),self.assertRaises(ValueError):normalize_record('BULK탱크DB',{'거래처코드':'C001','가스종류':'질소','탱크번호':'T001','소유구분':'당사임대',**bad},{'C001'})

 def test_reading_missing_tables_never_writes(self):
  service=fake_service({'통합업무':[['preserved']]})
  self.assertTrue(all(not r for r in CustomerStore(service,'sid').read_all().values()))
  service.spreadsheets.return_value.batchUpdate.assert_not_called()
  service.spreadsheets.return_value.values.return_value.batchUpdate.assert_not_called()

 def test_import_links_multiple_contacts_and_tanks_without_touching_tasks(self):
  tables={'통합업무':[['existing task']]};service=fake_service(tables);store=CustomerStore(service,'sid')
  incoming={'거래처DB':[customer()], '거래처연락처DB':[{'거래처코드':'C001','성명':'구매담당'},{'거래처코드':'C001','성명':'설비담당'}], 'BULK탱크DB':[{'거래처코드':'C001','가스종류':'질소','탱크번호':'T001','소유구분':'당사임대'},{'거래처코드':'C001','가스종류':'산소','탱크번호':'T002','소유구분':'당사임대'}]}
  self.assertEqual(store.import_new(incoming),5)
  loaded=store.read_all()
  self.assertEqual(len(loaded['BULK탱크DB']),2)
  self.assertEqual(len(loaded['거래처연락처DB']),2)
  self.assertEqual(tables['통합업무'],[['existing task']])
  service.spreadsheets.return_value.values.return_value.batchUpdate.reset_mock()
  with self.assertRaises(ValueError):store.import_new({'거래처DB':[customer()]})
  service.spreadsheets.return_value.values.return_value.batchUpdate.assert_not_called()

 def test_invalid_import_is_preflighted_before_writes(self):
  service=fake_service({});store=CustomerStore(service,'sid')
  with self.assertRaises(ValueError):store.import_new({'거래처DB':[customer()],'거래처연락처DB':[{'거래처코드':'OTHER','성명':'담당자'}]})
  service.spreadsheets.return_value.batchUpdate.assert_not_called()
  service.spreadsheets.return_value.values.return_value.batchUpdate.assert_not_called()

 def test_save_uses_physical_rows_and_protects_code(self):
  a,b=customer('C001'),customer('C002')
  headers=SCHEMAS['거래처DB']
  tables={'거래처DB':[headers,[a[k] for k in headers],[],[b[k] for k in headers]]}
  store=CustomerStore(fake_service(tables),'sid')
  b['거래처명']='변경 이름';store.save('거래처DB',b,'C002')
  self.assertEqual(tables['거래처DB'][3][1],'변경 이름')
  self.assertEqual(tables['거래처DB'][2],[])
  with self.assertRaises(ValueError):store.save('거래처DB',customer('C003'),'C002')

 def test_due_alerts_use_supplied_dates_and_exclude_recovered_tanks(self):
  alerts=upcoming_tanks([{'탱크번호':'T1','차기검사일':'2026-10-28','계약종료일':'2026-11-15','임대상태':'운영중'},{'차기검사일':'2026-10-28','임대상태':'회수완료'},{'차기검사일':'','임대상태':'운영중'}],date(2026,10,29))
  self.assertEqual(sorted(days for _,_,days in alerts),[-1,17])

 def test_empty_directory_registers_new_customer_in_dialog(self):
  import customer_workspace
  service=fake_service({});store=CustomerStore(service,'sid')
  from unittest.mock import patch
  with patch.object(customer_workspace,'_test_store',store,create=True):
   at=AppTest.from_string('from customer_workspace import render_customers,_test_store\nrender_customers(_test_store,["김주원"],lambda x:x)').run(timeout=30)
   self.assertFalse(at.exception)
   next(b for b in at.button if b.label=='거래처 등록').click().run(timeout=30)
   self.assertFalse(at.exception)
   next(x for x in at.text_input if x.label=='거래처명 *').set_value('새 거래처')
   next(x for x in at.selectbox if x.label=='거래구분 *').set_value('직거래처')
   next(b for b in at.button if b.label=='저장').click().run(timeout=30)
   self.assertFalse(at.exception)
   self.assertEqual(store.read_all()['거래처DB'][0]['거래처명'],'새 거래처')
   self.assertNotIn('customer_editor',at.session_state)

 def test_staff_directory_filters_by_owner_and_customer_name(self):
  import customer_workspace
  from unittest.mock import patch
  a=customer('A','가 거래처');a['당사담당자']='김주원'
  b=customer('B','나 거래처');b['당사담당자']='이상인'
  tables={'거래처DB':[SCHEMAS['거래처DB'],[a[k] for k in SCHEMAS['거래처DB']],[b[k] for k in SCHEMAS['거래처DB']]]}
  store=CustomerStore(fake_service(tables),'sid')
  with patch.object(customer_workspace,'_test_store',store,create=True):
   at=AppTest.from_string('from customer_workspace import render_customers,_test_store\nrender_customers(_test_store,["김주원","이상인"],lambda x:x)').run(timeout=30)
   self.assertFalse(at.exception)
   self.assertEqual(len(at.metric),0)
   at.selectbox(key='customer_staff_owner').set_value('이상인').run()
   self.assertTrue(any(b.key=='customer_staff_pick_B' for b in at.button))
   self.assertFalse(any(b.key=='customer_staff_pick_A' for b in at.button))
   at.text_input(key='customer_staff_search').set_value('없는 거래처').run()
   self.assertFalse(any((b.key or '').startswith('customer_staff_pick_') for b in at.button))
   self.assertFalse(at.exception)
