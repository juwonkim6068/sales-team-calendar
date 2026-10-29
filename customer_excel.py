"""Styled customer workbook and strictly validated import reader."""
from datetime import date
from io import BytesIO
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from customer_data import SCHEMAS, REQUIRED, CHOICES, DATE_FIELDS, text_value

DESCRIPTIONS={
 '거래처코드':'거래처를 연결하는 고유 코드. 예: C001. 연락처·탱크·특이사항에서 같은 코드를 사용합니다.',
 '거래처명':'정식 상호명', '거래구분':'최종 사용처인 직거래처 또는 재판매하는 대리점',
 '당사담당자':'당사 직원의 이름. 팀원정보의 이름과 동일하게 입력합니다.', '회사소개':'어떤 사업을 하는 회사인지, 규모·특징 등',
 '주요제품':'거래처가 생산·판매하는 제품', '취급가스':'당사가 공급하거나 거래처가 사용하는 가스. 예: 질소, 산소, 아르곤, 헬륨',
 '납품주소':'실제 가스 납품·탱크 설치 사업장 주소', '결제조건':'결제일, 외상 한도 및 수금 유의사항',
 '연락처':'직통번호 또는 휴대전화. 앞자리 0을 포함해 문자로 입력합니다.', '주담당여부':'주요 연락 창구이면 예. 한 거래처에 여러 담당자를 등록할 수 있습니다.',
 '용량':'숫자만 입력하고 단위는 용량단위 열에서 선택합니다.', '월임대료':'원 단위 숫자. 부가세·운송비 등 별도 조건은 특이사항에 기재합니다.',
 '탱크번호':'현장 설비 식별번호. 탱크 한 대당 한 행을 사용합니다.', '소유구분':'당사임대/거래처소유/타사임대 구분',
 '차기검사일':'안전관리자가 확인한 검사 예정일. 시스템이 법정 검사주기를 자동 산정하지 않습니다.',
 '설비담당자':'탱크 관리에 책임이 있는 당사 직원명', '중요도':'일반/중요/긴급', '내용':'사실·발생일·거래처 요청을 구체적으로 기록합니다.',
 '후속조치':'확인할 내용, 해야 할 일 및 조치 결과', '조치담당자':'후속조치를 맡은 당사 직원명',
 '직원명':'직원 이름. 팀원정보의 이름과 동일하게 입력합니다.', '담당업무':'담당 범위와 책임. 예: 특정 지역 가스 영업, BULK 납품 일정 관리',
 '협업자':'함께 수행하는 직원 이름', '대체담당자':'휴가·부재 시 대체하는 직원 이름', '수정일':'시스템에서 자동 기록합니다. 양식 작성 시 비워 두세요.',
}


def workbook_bytes():
 wb=Workbook();guide=wb.active;guide.title='작성안내'
 instructions=[
  ['거래처·BULK 설비·직원 담당업무 관리 양식',''],
  ['목적','거래처가 어떤 회사인지, 누가 담당하는지, 어떤 제품·가스를 취급하는지, 임대 탱크와 조치사항을 연결해 관리합니다.'],
  ['입력 순서','① 거래처DB → ② 거래처연락처DB/BULK탱크DB/거래처특이사항DB → ③ 직원담당업무DB'],
  ['거래처 연결','거래처코드(C001 등)를 먼저 정하고 관련 탭에서 동일하게 입력하세요. 거래처 한 곳에 여러 연락처·탱크·특이사항을 등록할 수 있습니다.'],
  ['필수 항목','데이터 탭에서 주황색으로 표시된 열은 필수입니다. 나머지는 필요한 만큼 작성하세요. 예시 업체·실제 고객 정보는 포함하지 않았습니다.'],
  ['자동 식별코드','연락처ID·탱크ID·특이사항ID·업무ID는 비워 두면 가져올 때 자동 생성됩니다. 거래처코드는 필수입니다.'],
  ['직원 이름','당사담당자·조치담당자·설비담당자·직원명은 기존 팀원정보의 이름으로 입력하세요. 직급은 대시보드에서 표시합니다.'],
  ['날짜·금액','날짜는 YYYY-MM-DD, 금액·용량은 쉼표·단위 없는 숫자로 입력합니다. 전화번호는 문자로 보존합니다.'],
  ['권장 기록','영업: 단가 조정 요청·경쟁 동향 / 공급: 납품시간·현장 출입 / 설비안전: 탱크 이상·검사예정 / 수금: 결제 지연 / 품질: 성적서·클레임'],
  ['BULK 탱크','한 대당 한 행. 탱크번호·가스·용량·소유구분·계약기간·검사일·설치위치를 기록합니다. 검사일은 안전관리자가 확인해 직접 입력합니다.'],
  ['대시보드 반영','거래처 특이사항 → 엑셀 가져오기에서 이 파일을 선택하고 검증 내용을 확인한 후 신규 추가합니다.'],
  ['기존 데이터 보호','엑셀 가져오기는 신규 추가 전용입니다. 기존 코드와 중복되면 중단합니다. 수정은 대시보드의 수정 버튼을 이용하세요.'],
  ['보호 대상','기존 업무일정·통합업무·메모·회의록·인증 설정은 변경하지 않습니다. 새 거래처 관리 탭만 사용합니다.'],
  ['수식 사용 금지','입력 탭에는 수식 대신 값만 입력하세요. 거래처명·연락처·가격 등은 내부 자료이므로 파일 접근 권한을 관리하세요.'],
 ]
 for row in instructions:guide.append(row)
 guide.column_dimensions['A'].width=23;guide.column_dimensions['B'].width=110
 guide.freeze_panes='B2'
 guide.row_dimensions[1].height=34
 guide.merge_cells('A1:B1')
 dictionary=wb.create_sheet('항목설명');dictionary.append(['시트','항목','필수','입력 가이드'])
 for sheet,headers in SCHEMAS.items():
  ws=wb.create_sheet(sheet);ws.append(headers)
  ws.freeze_panes='C2';ws.auto_filter.ref=f'A1:{ws.cell(101,len(headers)).column_letter}101'
  ws.sheet_view.zoomScale=85
  for col,field in enumerate(headers,1):
   cell=ws.cell(1,col)
   note=DESCRIPTIONS.get(field) or (' / '.join(CHOICES[field]) if field in CHOICES else ('날짜: YYYY-MM-DD' if field in DATE_FIELDS else '필요한 내용을 입력하세요.'))
   dictionary.append([sheet,field,'필수' if field in REQUIRED[sheet] else '선택',note])
   cell.comment=Comment(note,'입력 안내')
   ws.column_dimensions[cell.column_letter].width=38 if field in ('회사소개','내용','특이사항','담당업무','사업장주소','납품주소') else (25 if field in ('주요제품','취급가스','후속조치','결제조건') else 18)
   for row in range(2,102):
    target=ws.cell(row,col)
    target.number_format='@'
    target.alignment=Alignment(vertical='top',wrap_text=True)
    target.fill=PatternFill('solid',fgColor='F4F7FC' if row%2==0 else 'FFFFFF')
   if field in CHOICES:
    dv=DataValidation(type='list',formula1='"'+','.join(CHOICES[field])+'"',allow_blank=True)
    dv.errorTitle='목록에서 선택';dv.error='드롭다운의 값 중 하나를 선택하세요.';dv.showErrorMessage=True
    ws.add_data_validation(dv);dv.add(f'{cell.column_letter}2:{cell.column_letter}5001')
   if field in REQUIRED[sheet]:cell.fill=PatternFill('solid',fgColor='A65B10')
  ws.row_dimensions[1].height=32
  for row in range(2,102):ws.row_dimensions[row].height=30
 for ws in wb:
  for cell in ws[1]:
   if cell.fill.fgColor.rgb not in ('00A65B10','A65B10'):cell.fill=PatternFill('solid',fgColor='23476D')
   cell.font=Font(name='맑은 고딕',size=11,bold=True,color='FFFFFF')
   cell.alignment=Alignment(vertical='center',wrap_text=True)
  for row in ws.iter_rows(min_row=2):
   for cell in row:
    cell.font=Font(name='맑은 고딕',size=10,color='33465D')
    cell.alignment=Alignment(vertical='top',wrap_text=True)
  ws.sheet_properties.pageSetUpPr.fitToPage=True
  ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
  ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
 for row in range(2,len(instructions)+1):guide.row_dimensions[row].height=44
 for key,width in {'A':25,'B':24,'C':10,'D':100}.items():dictionary.column_dimensions[key].width=width
 dictionary.freeze_panes='D2';dictionary.auto_filter.ref=dictionary.dimensions
 buffer=BytesIO();wb.save(buffer);return buffer.getvalue()


def read_workbook(data):
 if len(data)>10*1024*1024:raise ValueError('엑셀 파일은 10MB 이하로 업로드해 주세요.')
 with ZipFile(BytesIO(data)) as archive:
  if sum(x.file_size for x in archive.infolist())>50*1024*1024:raise ValueError('압축 해제된 엑셀 데이터가 너무 큽니다.')
 wb=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
 try:
  missing=[s for s in SCHEMAS if s not in wb.sheetnames]
  if missing:raise ValueError('양식의 데이터 탭이 없습니다: '+', '.join(missing))
  result={}
  for sheet,headers in SCHEMAS.items():
   ws=wb[sheet]
   if ws.max_row>5001:raise ValueError(f'{sheet}: 한 번에 최대 5,000행만 가져올 수 있습니다.')
   iterator=ws.iter_rows()
   first=next(iterator,())
   if [text_value(c.value) for c in first]!=headers:raise ValueError(f'{sheet}: 열 이름·순서가 양식과 다릅니다.')
   rows=[]
   for cells in iterator:
    if any(c.data_type=='f' for c in cells):raise ValueError(f'{sheet}: 수식은 가져올 수 없습니다. 값으로 붙여넣으세요.')
    values=[text_value(c.value) for c in cells]
    if any(values):rows.append(dict(zip(headers,values)))
   result[sheet]=rows
  return result
 finally:wb.close()
